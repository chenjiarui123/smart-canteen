"""Safe CSV/XLSX parsing and deterministic data validation."""

from __future__ import annotations

import io
import math
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .context_builder import WEEKDAY_NAMES
from .dataset_repository import (
    SALES_COLUMNS,
    DatasetRepository,
    normalize_dish_name,
    safe_filename,
)


SALES_BASE_REQUIRED = {
    "餐次",
    "菜品名称",
    "菜品总销量",
    "菜品售卖单位",
    "订单数量",
    "菜品单价",
    "销售额",
    "平均取菜量",
}
WEATHER_COLUMNS = [
    "日期",
    "气温（℃）",
    "相对湿度（%）",
    "降水量（毫米）",
    "10米风速（公里/小时）",
]
CALENDAR_COLUMNS = [
    "日期",
    "星期",
    "是否校历寒暑假",
    "校历假期类型",
    "是否法定节假日假期",
    "法定节假日名称",
    "是否调休工作日",
]


@dataclass
class ParsedUpload:
    frame: pd.DataFrame | None
    sheet_names: list[str]
    requires_sheet_selection: bool
    safe_source_name: str


@dataclass
class PreviewArtifact:
    kind: str
    dataset_id: str | None
    processed: pd.DataFrame
    summary: dict[str, Any]
    rows: list[dict[str, Any]]
    issues: list[dict[str, Any]]
    can_commit: bool
    source_file: str


def _records(frame: pd.DataFrame, limit: int = 20) -> list[dict[str, Any]]:
    sample = frame.head(limit).copy()
    sample = sample.replace({np.nan: None, pd.NaT: None})
    for column in sample.columns:
        if pd.api.types.is_datetime64_any_dtype(sample[column]):
            sample[column] = sample[column].dt.strftime("%Y-%m-%d")
    result: list[dict[str, Any]] = []
    for row in sample.to_dict("records"):
        clean = {}
        for key, value in row.items():
            if isinstance(value, (np.integer,)):
                value = int(value)
            elif isinstance(value, (np.floating,)):
                value = float(value) if math.isfinite(float(value)) else None
            clean[str(key)] = value
        result.append(clean)
    return result


def _read_csv_bytes(content: bytes) -> pd.DataFrame:
    last_error: Exception | None = None
    for encoding in ["utf-8-sig", "utf-8"]:
        try:
            return pd.read_csv(io.BytesIO(content), encoding=encoding)
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError("CSV 必须使用 UTF-8 或 UTF-8-SIG 编码") from last_error


def _xlsx_sheet_names(content: bytes) -> list[str]:
    workbook = load_workbook(
        io.BytesIO(content),
        read_only=True,
        data_only=False,
        keep_links=False,
    )
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _reject_formulas(content: bytes, sheet_name: str) -> None:
    workbook = load_workbook(
        io.BytesIO(content),
        read_only=True,
        data_only=False,
        keep_links=False,
    )
    try:
        worksheet = workbook[sheet_name]
        for row in worksheet.iter_rows():
            if any(cell.data_type == "f" for cell in row):
                raise ValueError("上传工作表包含公式；请转换为静态数值后再上传")
    finally:
        workbook.close()


def parse_upload(
    filename: str, content: bytes, sheet_name: str | None
) -> ParsedUpload:
    safe_name = safe_filename(filename)
    suffix = Path(safe_name).suffix.lower()
    if suffix == ".csv":
        return ParsedUpload(
            frame=_read_csv_bytes(content),
            sheet_names=[],
            requires_sheet_selection=False,
            safe_source_name=safe_name,
        )
    if suffix != ".xlsx":
        raise ValueError("仅支持 UTF-8 CSV 或不含宏的 XLSX 文件")
    sheets = _xlsx_sheet_names(content)
    if not sheets:
        raise ValueError("Excel 文件没有可读取工作表")
    if sheet_name is None and len(sheets) > 1:
        return ParsedUpload(
            frame=None,
            sheet_names=sheets,
            requires_sheet_selection=True,
            safe_source_name=safe_name,
        )
    selected = sheet_name or sheets[0]
    if selected not in sheets:
        raise ValueError("所选工作表不存在")
    _reject_formulas(content, selected)
    frame = pd.read_excel(
        io.BytesIO(content),
        sheet_name=selected,
        engine="openpyxl",
    )
    return ParsedUpload(
        frame=frame,
        sheet_names=sheets,
        requires_sheet_selection=False,
        safe_source_name=safe_name,
    )


def _normalise_flag(value: Any) -> int | None:
    if pd.isna(value):
        return None
    text = str(value).strip().lower()
    if text in {"1", "true", "yes", "是", "工作日"}:
        return 1
    if text in {"0", "false", "no", "否", "非工作日"}:
        return 0
    return None


def _has_calendar_label(value: Any, empty_labels: set[str]) -> bool:
    text = str(value or "").strip().lower()
    return text not in empty_labels


def _calendar_label_issues(work: pd.DataFrame) -> pd.Series:
    non_holiday_labels = {"", "nan", "none", "非节假日", "非假期", "无"}
    non_break_labels = non_holiday_labels | {"非寒暑假"}
    holiday_flag = work["是否法定节假日假期"]
    break_flag = work["是否校历寒暑假"]
    holiday_name_present = work["法定节假日名称"].map(
        lambda value: _has_calendar_label(value, non_holiday_labels)
    )
    break_type_present = work["校历假期类型"].map(
        lambda value: _has_calendar_label(value, non_break_labels)
    )
    holiday_conflict = holiday_flag.notna() & holiday_flag.eq(1).ne(
        holiday_name_present
    )
    break_conflict = break_flag.notna() & break_flag.eq(1).ne(
        break_type_present
    )
    return holiday_conflict | break_conflict


def _dates_from_frame(frame: pd.DataFrame) -> pd.Series:
    if "日期" in frame.columns:
        return pd.to_datetime(frame["日期"], errors="coerce")
    if {"年", "月", "日"}.issubset(frame.columns):
        return pd.to_datetime(
            {
                "year": pd.to_numeric(frame["年"], errors="coerce"),
                "month": pd.to_numeric(frame["月"], errors="coerce"),
                "day": pd.to_numeric(frame["日"], errors="coerce"),
            },
            errors="coerce",
        )
    return pd.Series(pd.NaT, index=frame.index, dtype="datetime64[ns]")


def _issue(
    issues: list[dict[str, Any]],
    issue_type: str,
    message: str,
    rows: list[int] | None = None,
) -> None:
    issues.append(
        {
            "type": issue_type,
            "message": message,
            "rows": rows or [],
        }
    )


def preprocess_sales(
    frame: pd.DataFrame,
    dataset_id: str,
    source_file: str,
    repository: DatasetRepository,
) -> PreviewArtifact:
    original_rows = int(len(frame))
    issues: list[dict[str, Any]] = []
    missing = sorted(SALES_BASE_REQUIRED - set(frame.columns))
    if missing:
        _issue(issues, "missing_fields", f"缺少字段：{', '.join(missing)}")
        return PreviewArtifact(
            "sales",
            dataset_id,
            pd.DataFrame(columns=SALES_COLUMNS),
            {
                "original_rows": original_rows,
                "valid_rows": 0,
                "duplicate_rows": 0,
                "abnormal_rows": original_rows,
            },
            _records(frame),
            issues,
            False,
            source_file,
        )

    work = frame.copy()
    work["_date"] = _dates_from_frame(work)
    invalid = work["_date"].isna()
    if invalid.any():
        _issue(
            issues,
            "invalid_date",
            "存在无法解析的日期",
            (work.index[invalid] + 2).astype(int).tolist()[:100],
        )

    meal_map = {
        "早": "早餐",
        "早饭": "早餐",
        "早餐": "早餐",
        "中": "午餐",
        "中餐": "午餐",
        "午餐": "午餐",
        "晚": "晚餐",
        "晚饭": "晚餐",
        "晚餐": "晚餐",
    }
    work["餐次"] = work["餐次"].astype(str).str.strip().map(meal_map)
    bad_meal = work["餐次"].isna()
    invalid |= bad_meal
    if bad_meal.any():
        _issue(
            issues,
            "invalid_meal",
            "餐次必须标准化为早餐、午餐或晚餐",
            (work.index[bad_meal] + 2).astype(int).tolist()[:100],
        )

    numeric_columns = ["菜品总销量", "订单数量", "销售额", "平均取菜量"]
    for column in numeric_columns:
        work[column] = pd.to_numeric(work[column], errors="coerce")
        bad = (
            work[column].isna()
            | ~np.isfinite(work[column])
            | work[column].lt(0)
        )
        invalid |= bad
        if bad.any():
            _issue(
                issues,
                "invalid_numeric",
                f"{column} 必须为有限非负数",
                (work.index[bad] + 2).astype(int).tolist()[:100],
            )
    empty_dish = work["菜品名称"].isna() | work["菜品名称"].astype(str).str.strip().eq("")
    invalid |= empty_dish
    if empty_dish.any():
        _issue(issues, "empty_dish", "菜品名称不能为空")
    work["菜品名称"] = work["菜品名称"].astype(str).str.strip()
    work["_dish_norm"] = work["菜品名称"].map(normalize_dish_name)

    date_text = work["_date"].dt.strftime("%Y-%m-%d")
    duplicate_in_file = pd.Series(
        list(zip(date_text, work["餐次"], work["_dish_norm"])),
        index=work.index,
    ).duplicated(keep="first")
    if duplicate_in_file.any():
        _issue(
            issues,
            "duplicate_in_upload",
            "上传文件内存在重复业务键，重复行不会写入",
            (work.index[duplicate_in_file] + 2).astype(int).tolist()[:100],
        )

    existing = repository.combined_sales_raw(dataset_id)
    existing_dates = pd.to_datetime(
        {
            "year": existing["年"],
            "month": existing["月"],
            "day": existing["日"],
        }
    )
    existing_keys = set(
        zip(
            existing_dates.dt.strftime("%Y-%m-%d"),
            existing["餐次"].astype(str),
            existing["菜品名称"].map(normalize_dish_name),
        )
    )
    existing_duplicate = pd.Series(
        [
            key in existing_keys
            for key in zip(date_text, work["餐次"], work["_dish_norm"])
        ],
        index=work.index,
    )
    duplicate = duplicate_in_file | existing_duplicate

    weather = repository.read_weather().copy()
    calendar = repository.read_calendar().copy()
    weather_dates = set(weather["日期"].astype(str))
    calendar_dates = set(calendar["日期"].astype(str))
    upload_dates = set(date_text.dropna().astype(str))
    missing_weather = sorted(upload_dates - weather_dates)
    missing_calendar = sorted(upload_dates - calendar_dates)
    if missing_weather:
        _issue(
            issues,
            "missing_weather",
            "以下日期缺少天气，必须先更新天气数据",
        )
    if missing_calendar:
        _issue(
            issues,
            "missing_calendar",
            "以下日期缺少校历，必须先更新校历数据",
        )

    dataset = repository.dataset(dataset_id)
    context = weather.merge(calendar, on="日期", how="inner")
    context["_date"] = pd.to_datetime(context["日期"], errors="raise")
    context = context.drop(columns="日期").drop_duplicates("_date")
    work = work.drop(
        columns=[
            column
            for column in [
                "星期",
                "是否校历寒暑假",
                "校历假期类型",
                "是否法定节假日假期",
                "法定节假日名称",
                "是否调休工作日",
                "气温（℃）",
                "相对湿度（%）",
                "降水量（毫米）",
                "10米风速（公里/小时）",
            ]
            if column in work.columns
        ]
    )
    work = work.merge(context, on="_date", how="left")
    work["年"] = work["_date"].dt.year
    work["月"] = work["_date"].dt.month
    work["日"] = work["_date"].dt.day
    work["食堂名称"] = dataset["display_name"]

    library = repository.read_dish_library(dataset_id)
    library_names = set(library["菜品名称"].map(normalize_dish_name))
    unmatched = sorted(
        set(work.loc[~invalid, "_dish_norm"]) - library_names
    )
    if unmatched:
        display = (
            work.loc[work["_dish_norm"].isin(unmatched), "菜品名称"]
            .drop_duplicates()
            .sort_values()
            .tolist()
        )
        _issue(
            issues,
            "unmatched_dishes",
            f"{len(display)} 道菜未匹配菜品库，将仅使用保守菜名推断",
        )
    existing_dishes = set(existing["菜品名称"].map(normalize_dish_name))
    new_dishes = sorted(set(work.loc[~invalid, "_dish_norm"]) - existing_dishes)
    valid_mask = ~invalid & ~duplicate
    processed = work.loc[valid_mask].copy()
    for column in SALES_COLUMNS:
        if column not in processed:
            processed[column] = ""
    processed = processed[SALES_COLUMNS]
    date_range = [
        date_text.dropna().min(),
        date_text.dropna().max(),
    ] if date_text.notna().any() else [None, None]
    summary = {
        "original_rows": original_rows,
        "valid_rows": int(valid_mask.sum()),
        "duplicate_rows": int(duplicate.sum()),
        "abnormal_rows": int(invalid.sum()),
        "date_range": date_range,
        "new_dish_count": len(new_dishes),
        "unmatched_dish_count": len(unmatched),
        "unmatched_dishes": (
            work.loc[work["_dish_norm"].isin(unmatched), "菜品名称"]
            .drop_duplicates()
            .sort_values()
            .tolist()
        ),
        "missing_weather_dates": missing_weather,
        "missing_calendar_dates": missing_calendar,
    }
    can_commit = bool(
        len(processed) > 0 and not missing_weather and not missing_calendar
    )
    return PreviewArtifact(
        "sales",
        dataset_id,
        processed,
        summary,
        _records(work.drop(columns=["_dish_norm"], errors="ignore")),
        issues,
        can_commit,
        source_file,
    )


def preprocess_weather(
    frame: pd.DataFrame,
    source_file: str,
    repository: DatasetRepository,
) -> PreviewArtifact:
    issues: list[dict[str, Any]] = []
    missing = sorted(set(WEATHER_COLUMNS) - set(frame.columns))
    if missing:
        _issue(issues, "missing_fields", f"缺少字段：{', '.join(missing)}")
        return PreviewArtifact(
            "weather",
            None,
            pd.DataFrame(columns=WEATHER_COLUMNS),
            {"original_rows": len(frame), "new_dates": [], "existing_dates": [], "conflict_dates": []},
            _records(frame),
            issues,
            False,
            source_file,
        )
    work = frame[WEATHER_COLUMNS].copy()
    parsed = pd.to_datetime(work["日期"], errors="coerce")
    invalid = parsed.isna()
    work["日期"] = parsed.dt.strftime("%Y-%m-%d")
    option_map = repository.ui_options()["weather"]
    field_map = {
        "气温（℃）": "temperature_c",
        "相对湿度（%）": "humidity_pct",
        "降水量（毫米）": "precipitation_mm",
        "10米风速（公里/小时）": "wind_speed_kmh",
    }
    for column, key in field_map.items():
        work[column] = pd.to_numeric(work[column], errors="coerce")
        values = option_map[key]["values"]
        bad = work[column].isna() | ~work[column].round(1).isin(values)
        invalid |= bad
        if bad.any():
            _issue(
                issues,
                "invalid_weather_value",
                f"{column} 存在超出允许下拉范围的数值",
                (work.index[bad] + 2).astype(int).tolist()[:100],
            )
    duplicate = work["日期"].duplicated(keep="first")
    current = repository.read_weather()
    current_lookup = current.set_index("日期").to_dict("index")
    existing_dates: list[str] = []
    conflict_dates: list[str] = []
    new_dates: list[str] = []
    for row in work.loc[~invalid & ~duplicate].to_dict("records"):
        current_row = current_lookup.get(row["日期"])
        if current_row is None:
            new_dates.append(row["日期"])
        elif all(
            abs(float(current_row[column]) - float(row[column])) <= 1e-9
            for column in WEATHER_COLUMNS[1:]
        ):
            existing_dates.append(row["日期"])
        else:
            conflict_dates.append(row["日期"])
    if conflict_dates:
        _issue(
            issues,
            "weather_conflict",
            "冲突日期不会覆盖已有天气记录",
        )
    processed = work.loc[
        ~invalid & ~duplicate & work["日期"].isin(new_dates)
    ].copy()
    summary = {
        "original_rows": int(len(frame)),
        "valid_rows": int((~invalid).sum()),
        "duplicate_rows": int(duplicate.sum()),
        "abnormal_rows": int(invalid.sum()),
        "new_dates": sorted(new_dates),
        "existing_dates": sorted(existing_dates),
        "conflict_dates": sorted(conflict_dates),
        "missing_fields": [],
    }
    return PreviewArtifact(
        "weather",
        None,
        processed,
        summary,
        _records(work),
        issues,
        bool(len(processed)),
        source_file,
    )


def preprocess_calendar(
    frame: pd.DataFrame,
    source_file: str,
    repository: DatasetRepository,
) -> PreviewArtifact:
    issues: list[dict[str, Any]] = []
    missing = sorted(set(CALENDAR_COLUMNS) - set(frame.columns))
    if missing:
        _issue(issues, "missing_fields", f"缺少字段：{', '.join(missing)}")
        return PreviewArtifact(
            "calendar",
            None,
            pd.DataFrame(columns=CALENDAR_COLUMNS),
            {"original_rows": len(frame), "new_dates": [], "existing_dates": [], "conflict_dates": []},
            _records(frame),
            issues,
            False,
            source_file,
        )
    work = frame[CALENDAR_COLUMNS].copy()
    parsed = pd.to_datetime(work["日期"], errors="coerce")
    invalid = parsed.isna()
    work["日期"] = parsed.dt.strftime("%Y-%m-%d")
    expected_weekday = parsed.map(
        lambda value: WEEKDAY_NAMES[value.weekday()] if pd.notna(value) else None
    )
    bad_weekday = work["星期"].astype(str).ne(expected_weekday)
    invalid |= bad_weekday
    if bad_weekday.any():
        _issue(issues, "invalid_weekday", "星期与日期计算结果不一致")
    for column in [
        "是否校历寒暑假",
        "是否法定节假日假期",
        "是否调休工作日",
    ]:
        flags = work[column].map(_normalise_flag)
        bad = flags.isna()
        invalid |= bad
        work[column] = flags
        if bad.any():
            _issue(issues, "invalid_flag", f"{column} 必须为是/否或1/0")
    label_conflicts = _calendar_label_issues(work)
    invalid |= label_conflicts
    if label_conflicts.any():
        _issue(
            issues,
            "calendar_flag_label_conflict",
            "校历假期标志与假期类型或节假日名称不一致",
            (work.index[label_conflicts] + 2).astype(int).tolist()[:100],
        )
    duplicate = work["日期"].duplicated(keep="first")
    current = repository.read_calendar()
    current_lookup = current.set_index("日期").astype(str).to_dict("index")
    existing_dates: list[str] = []
    conflict_dates: list[str] = []
    new_dates: list[str] = []
    for row in work.loc[~invalid & ~duplicate].to_dict("records"):
        current_row = current_lookup.get(row["日期"])
        if current_row is None:
            new_dates.append(row["日期"])
        elif all(
            str(current_row[column]) == str(row[column])
            for column in CALENDAR_COLUMNS[1:]
        ):
            existing_dates.append(row["日期"])
        else:
            conflict_dates.append(row["日期"])
    if conflict_dates:
        _issue(
            issues,
            "calendar_conflict",
            "冲突日期不会覆盖已有校历记录",
        )
    processed = work.loc[
        ~invalid & ~duplicate & work["日期"].isin(new_dates)
    ].copy()
    summary = {
        "original_rows": int(len(frame)),
        "valid_rows": int((~invalid).sum()),
        "duplicate_rows": int(duplicate.sum()),
        "abnormal_rows": int(invalid.sum()),
        "new_dates": sorted(new_dates),
        "existing_dates": sorted(existing_dates),
        "conflict_dates": sorted(conflict_dates),
        "missing_fields": [],
    }
    return PreviewArtifact(
        "calendar",
        None,
        processed,
        summary,
        _records(work),
        issues,
        bool(len(processed)),
        source_file,
    )

from __future__ import annotations

from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "智慧食堂项目阶段报告_甲方汇报_20260824.docx"
SCREEN_RECO = ROOT / "outputs" / "client_report_web_recommendation_crop.png"
SCREEN_UPDATE = ROOT / "outputs" / "client_report_web_update_crop.png"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5E6B78"
LIGHT = "F2F4F7"
PALE_BLUE = "E8EEF5"
PALE_GREEN = "EDF5EE"
GREEN = "2F6B3A"
GOLD = "A97813"
RED = "9B1C1C"
WHITE = "FFFFFF"


def set_font(run, size=11, bold=False, color="000000", italic=False, name="Arial Unicode MS"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_text(cell, text, *, bold=False, color="000000", size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    p.clear()
    r = p.add_run(str(text))
    set_font(r, size=size, bold=bold, color=color)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_font(run, size=9, color=MUTED)


def add_header_footer(section, first=False):
    section.different_first_page_header_footer = first
    if first:
        section.first_page_header.paragraphs[0].text = ""
        section.first_page_footer.paragraphs[0].text = ""
    header = section.header
    p = header.paragraphs[0]
    p.text = "智慧食堂菜单推荐系统｜项目阶段总结"
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    set_font(p.runs[0], size=9, bold=True, color=MUTED)
    p_pr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "D8DEE6")
    pbdr.append(bottom)
    p_pr.append(pbdr)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = "项目阶段汇报材料  ·  "
    set_font(fp.runs[0], size=9, color=MUTED)
    add_page_field(fp)


def add_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_id = 100
    num_id = 100
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl.append(lvl_text)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(tabs)
    ppr.append(ind)
    lvl.append(ppr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_bullet(doc, text, num_id):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(numid)
    p_pr.append(num_pr)
    r = p.add_run(text)
    set_font(r, size=10.5, color="25313C")
    return p


def add_body(doc, text, *, bold_lead=None, color="25313C", after=6, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, size=11, bold=True, color=INK)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2, size=11, color=color)
    else:
        r = p.add_run(text)
        set_font(r, size=11, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    if level == 1 and getattr(doc, "_pending_page_break", False):
        p.paragraph_format.page_break_before = True
        doc._pending_page_break = False
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    p.add_run(text)
    return p


def add_kpi_strip(doc, items):
    table = doc.add_table(rows=1, cols=len(items))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    widths = [9360 // len(items)] * len(items)
    set_table_geometry(table, widths)
    for cell, (value, label) in zip(table.rows[0].cells, items):
        set_cell_shading(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(value)
        set_font(r, size=18, bold=True, color=DARK_BLUE)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(5)
        r2 = p2.add_run(label)
        set_font(r2, size=8.5, color=MUTED)
    return table


def add_callout(doc, title, body, fill=PALE_GREEN, accent=GREEN):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_font(r, size=11, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    p2.paragraph_format.line_spacing = 1.10
    r2 = p2.add_run(body)
    set_font(r2, size=10.5, color="25313C")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_table(doc, headers, rows, widths, numeric_cols=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, widths)
    for i, h in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], LIGHT)
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=INK, size=9.2,
                      align=WD_ALIGN_PARAGRAPH.CENTER if i in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=9.1,
                          align=WD_ALIGN_PARAGRAPH.CENTER if i in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc, path, caption):
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(path), width=Inches(6.15))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    r = cp.add_run(caption)
    set_font(r, size=9, color=MUTED, italic=True)


def page_break(doc):
    # Main sections flow continuously to avoid accidental blank pages when the
    # previous section already fills a page.  Only the cover requests a forced
    # break explicitly below.
    return None


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    add_header_footer(sec, first=True)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for level, size, color, before, after in [
        (1, 16, BLUE, 16, 8),
        (2, 13, BLUE, 12, 6),
        (3, 12, DARK_BLUE, 8, 4),
    ]:
        st = styles[f"Heading {level}"]
        st.font.name = "Arial Unicode MS"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    num_id = add_numbering(doc)

    # Cover — editorial_cover pattern
    doc.add_paragraph().paragraph_format.space_after = Pt(72)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    set_font(kicker.add_run("项目阶段汇报"), size=11, bold=True, color=GOLD)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    set_font(title.add_run("智慧食堂菜单推荐系统"), size=29, bold=True, color=INK)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(30)
    set_font(sub.add_run("项目报告与阶段总结"), size=17, color=DARK_BLUE)
    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = rule.add_run("━━━━━━━━━━━━━━━━━━━━")
    set_font(rr, size=10, color="9FB7C9")
    doc.add_paragraph().paragraph_format.space_after = Pt(55)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(8)
    set_font(meta.add_run("汇报对象：项目甲方及业务管理人员"), size=11, bold=True, color=INK)
    meta2 = doc.add_paragraph()
    meta2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta2.paragraph_format.space_after = Pt(8)
    set_font(meta2.add_run("材料日期：2026年8月24日"), size=10.5, color=MUTED)
    meta3 = doc.add_paragraph()
    meta3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(meta3.add_run("当前阶段：已完成模型、网页、数据更新与两套食堂数据接入"), size=10.5, color=MUTED)
    doc._pending_page_break = True

    add_heading(doc, "一、项目摘要", 1)
    add_callout(doc, "一句话结论", "项目已从“分析模型”推进为可在本地网页使用的菜单辅助决策系统。系统能够根据历史销售、日期、天气和假期信息，给出常用菜品、低频高销量潜力菜及建议销量，并允许工作人员人工选择、检查搭配和保存菜单。当前建议进入小范围试运行，而不是完全无人干预地自动排菜。")
    add_kpi_strip(doc, [
        ("2 套", "独立食堂数据体系"),
        ("21.1 万+", "历史销售记录"),
        ("4,345", "两套菜品库记录"),
        ("本地运行", "业务数据不外发"),
    ])
    add_heading(doc, "已经完成的核心成果", 2)
    for t in [
        "形成可解释的菜单推荐方法：先预测当餐总需求，再选择菜品并分配建议销量。",
        "建成网页端菜单推荐功能：输入食堂、日期、餐次、天气和假期信息即可生成候选菜单。",
        "建立双通道推荐：常用菜保证稳定供应，潜力菜补充低频但历史单次销量较高的选择。",
        "增加人工选择缓冲区、类别/销量筛选、菜单搭配提醒和可选的总量重新分配。",
        "建成数据更新页面：支持 CSV/XLSX 预览、校验、备份后再写入，原始历史数据保持只读。",
        "实现 Windows、macOS、Linux 本地运行，并对两套逻辑食堂数据进行隔离管理。",
    ]:
        add_bullet(doc, t, num_id)
    add_heading(doc, "对甲方的直接价值", 2)
    add_table(doc, ["业务环节", "当前改进"], [
        ["排菜单", "从完全依赖经验，变为“数据建议 + 人工确认”"],
        ["备餐量", "为每道候选菜给出建议销量，并提示总备餐量是否偏高或偏低"],
        ["菜品多样性", "除常用菜外，补充低频高销量潜力菜，并支持半随机刷新"],
        ["特殊日期", "把周末、法定节假日、寒暑假、调休和天气作为额外参考"],
        ["数据维护", "新增销售、天气和校历数据先预览、再确认、可备份恢复"],
    ], [1800, 7560])
    page_break(doc)

    add_heading(doc, "二、项目目标与建设范围", 1)
    add_body(doc, "项目要解决的不是“让系统完全替代食堂管理人员”，而是为每餐排菜提供一份可解释、可调整、可保存的数据建议，减少重复劳动，提高备餐判断的一致性。")
    add_heading(doc, "业务目标", 2)
    for t in [
        "根据预测日期以前的历史销售信息，判断当餐大致需要准备多少。",
        "在可供应菜品中推荐较合适的组合，并说明推荐原因。",
        "兼顾常用畅销菜、低频潜力菜、菜品搭配和节假日变化。",
        "让工作人员保留最终决定权，系统只做建议、提醒和记录。",
    ]:
        add_bullet(doc, t, num_id)
    add_heading(doc, "当前交付范围", 2)
    add_table(doc, ["模块", "已完成内容", "当前状态"], [
        ["菜单预测", "按食堂、日期、餐次生成菜品和建议销量", "已完成"],
        ["潜力菜推荐", "按类别和数量推荐低频高销量菜，可刷新", "已完成"],
        ["菜单确认", "已选菜品独立保留，刷新不丢失；可删除或追加", "已完成"],
        ["搭配检查", "检查重复菜、相似菜、食材/做法/口味集中等问题", "已完成"],
        ["推荐留档", "按日期和餐次保存 CSV 与推荐说明", "已完成"],
        ["数据更新", "销售、天气、校历上传、预览、校验、备份和写入", "已完成"],
        ["数据体系", "A湖畔、F基础园两套独立历史和菜品库", "已完成"],
    ], [1500, 6260, 1600], numeric_cols=(2,))
    add_callout(doc, "当前定位", "系统已经具备试运行条件，适合在真实排菜单流程中作为“第二意见”和记录工具。尚不建议直接取消人工审核。", fill="FFF8E8", accent=GOLD)
    page_break(doc)

    add_heading(doc, "三、数据基础与使用边界", 1)
    add_body(doc, "系统将原始历史销售与网页后续新增的真实销售分开保存。推荐时只读取预测日期以前的数据，避免把当天或未来真实销量带入预测。")
    add_table(doc, ["数据体系", "销售记录", "日期范围", "菜品库", "餐次"], [
        ["A湖畔智慧餐厅", "91,971 条", "2025-07-21 至 2026-06-10", "2,148 道", "早/午/晚"],
        ["F基础园智慧餐厅", "119,065 条", "2025-01-01 至 2026-06-10", "2,197 道", "早/午/晚"],
    ], [1800, 1500, 2700, 1500, 1860], numeric_cols=(1, 3, 4))
    add_heading(doc, "系统使用的信息", 2)
    for t in [
        "历史销售：菜品销量、订单数、销售额、价格、平均取菜量和历史出菜情况。",
        "时间信息：星期、普通工作日/休息日、调休、校历寒暑假、法定节假日。",
        "天气信息：气温、湿度、降水量和风速；预测当餐时允许使用用户输入的实际条件。",
        "菜品库：菜品类别、口味、颜色、营养，以及已维护的主要食材、做法、主食和相似菜信息。",
        "推荐历史：只记录工作人员最终确认的菜单，用于后续减少重复推荐。",
    ]:
        add_bullet(doc, t, num_id)
    add_heading(doc, "信息边界", 2)
    add_callout(doc, "严格时间规则", "预测某一天时，系统只使用该日期以前的销售记录；目标日的日期、天气和假期条件可以作为输入，但不会读取目标日的真实菜单、销量或营业额。", fill=PALE_BLUE, accent=DARK_BLUE)
    add_body(doc, "系统不联网搜索菜品答案，也不把业务数据发送到外部服务。菜品得分和销量来自固定代码与本地数据，不使用大模型进行主观打分。")
    page_break(doc)

    add_heading(doc, "四、用户如何使用系统", 1)
    add_figure(doc, SCREEN_RECO, "图1  菜单推荐页面：通过下拉框和日期控件输入用餐条件")
    add_table(doc, ["步骤", "工作人员操作", "系统输出"], [
        ["1", "选择食堂、日期、餐次、天气和假期情况", "建立本次预测条件"],
        ["2", "点击生成推荐菜单", "常用菜品与潜力菜两个候选区"],
        ["3", "按类别、预计销量筛选，刷新潜力菜", "更聚焦地查看候选菜"],
        ["4", "把需要的菜加入已选缓冲区", "已选菜不因筛选或刷新消失"],
        ["5", "查看总菜品数、总备餐量和搭配提醒", "数量、风险与搭配建议"],
        ["6", "确认并保存", "CSV 菜单、说明报告和推荐历史"],
    ], [720, 4320, 4320], numeric_cols=(0,))
    page_break(doc)

    add_heading(doc, "五、当前方法：先估总量，再选菜，再分配", 1)
    add_body(doc, "为了便于理解，可以把当前方法看成三步：先判断这餐大概卖多少，再判断哪些菜更值得进入候选，最后把总需求分配到具体菜品。")
    add_table(doc, ["步骤", "系统做什么", "通俗解释"], [
        ["估算当餐总需求", "结合近期与长期销量、星期、趋势、假期和天气", "先确定这餐大概需要准备多少"],
        ["确定菜单数量", "参考同类日期的真实菜品数和平均每菜销量", "特殊日期不应照搬普通工作日的菜单规模"],
        ["筛选候选菜", "看历史出现频率、近期活跃、销量稳定性和销售表现", "优先选择有证据支持的菜"],
        ["检查菜单搭配", "排除重复菜和明确冲突，其他相近项给提醒或轻度降分", "避免整餐过于相似"],
        ["分配建议销量", "按菜品权重和稳定性分配，并守住当餐总量", "各菜建议量之和与总需求保持一致"],
        ["剔除过小批量", "历史样本支持时，移除不值得单独备餐的极小数量", "避免出现只建议几份却增加生产复杂度的菜"],
    ], [1550, 4050, 3760])
    add_heading(doc, "近期信息与全部历史如何配合", 2)
    add_body(doc, "系统不会只看最近几天，也不会只看长期平均。近期数据反映最新变化，全部历史用于在样本不足时提供稳定参考。预测日期离历史数据很远时，系统会自动降低近期信息的权重，并提示“远期预测、近期特征不足”。")
    add_heading(doc, "特殊日期和天气如何使用", 2)
    for t in [
        "法定节假日、校历假期重叠时，只形成一个联合状态，避免两个下降因素重复叠加。",
        "调休工作日按工作日规律处理，但仍保留自身标记。",
        "天气只作为补充修正，假期已经明显改变需求时，天气影响会被弱化。",
        "所有下降或上升因素合并后设置合理上下限，防止多个因素相乘后把预测压得过低或抬得过高。",
    ]:
        add_bullet(doc, t, num_id)
    page_break(doc)

    add_heading(doc, "六、双通道菜单：稳定供应与潜力探索", 1)
    add_table(doc, ["推荐通道", "主要作用", "当前规则"], [
        ["常用菜品", "保证菜单稳定、覆盖历史常卖菜", "沿用冻结 Trial 18 策略，按历史频率、预测销量、稳定性和搭配规则推荐"],
        ["低频高销量潜力菜", "增加新鲜感，发现不常出现但单次表现好的菜", "独立策略文件；可按类别和数量生成，支持半随机刷新，不占常用菜名额"],
    ], [1900, 3000, 4460])
    add_heading(doc, "潜力菜推荐的关键设计", 2)
    for t in [
        "可选择“全部、荤菜、素菜、主食、甜品、水果”，并输入本次希望查看的数量。",
        "每次刷新会在合格候选中进行可复现的半随机变化，避免每次只看到固定几道菜。",
        "已经加入最终菜单的潜力菜会逐步降低后续推荐权重，但不设置绝对冷却。",
        "汤圆、粽子、月饼等节日菜在非节日也可推荐，但非节日评分不使用节日期间的销量高峰。",
        "潜力菜是额外的试餐建议，工作人员选入后再与常用菜一起检查总数量、总备餐量和整体搭配。",
    ]:
        add_bullet(doc, t, num_id)
    add_callout(doc, "为什么保留两栏", "常用菜解决“今天大概率要卖什么”，潜力菜解决“还能尝试什么”。两栏分开，既不破坏稳定菜单，也给管理人员留下探索空间。", fill=PALE_GREEN, accent=GREEN)
    page_break(doc)

    add_heading(doc, "七、菜单搭配与数量提醒", 1)
    add_body(doc, "系统把必须避免的问题与可以提醒的问题分开处理。这样既防止明显错误，又不会因为规则过严而把菜单压得过少。")
    add_table(doc, ["规则类型", "例子", "系统处理"], [
        ["必须遵守", "重复菜名、明确相似菜、高置信度同核心食材冲突", "阻止重复或冲突菜同时进入最终菜单"],
        ["建议优化", "相同做法过多、口味过于集中、同食材大类偏多、近期重复", "给出提醒或降低评分，不轻易一刀切"],
        ["数量参考", "总菜品数偏少/偏多、总备餐量偏离当餐需求", "给出意见，不强制修改；可选择按总需求重新分配"],
    ], [1650, 3740, 3970])
    add_heading(doc, "最终菜单确认时重新检查", 2)
    for t in [
        "重新统计常用菜与潜力菜合并后的总菜品数量。",
        "重新计算总备餐量，并判断是否高于或低于当餐参考需求。",
        "重新检查菜品类别、食材、做法、口味、相似菜和重复情况。",
        "用户可选择保留各菜原建议量，也可按当餐总需求重新分配。",
        "确认后才写入推荐历史；仅生成或刷新候选不会影响后续推荐。",
    ]:
        add_bullet(doc, t, num_id)
    page_break(doc)

    add_heading(doc, "八、数据更新与运行安全", 1)
    add_figure(doc, SCREEN_UPDATE, "图2  数据更新页面：上传后先预览，确认后才写入")
    add_heading(doc, "更新流程", 2)
    add_table(doc, ["环节", "保护措施"], [
        ["文件上传", "支持 CSV/XLSX；拒绝旧版 XLS、宏和公式"],
        ["预览校验", "显示有效行、重复、异常、日期范围、缺失天气/校历和未匹配菜品"],
        ["确认写入", "只有用户明确点击确认后才更新正式运行数据"],
        ["重复保护", "按日期、餐次和标准菜名识别；同一文件重复上传不会重复写入"],
        ["备份恢复", "写入前自动备份；失败时恢复；提供跨平台恢复脚本"],
        ["并发保护", "推荐保存与数据更新共用文件锁，避免同时写入造成损坏"],
    ], [1800, 7560])
    add_callout(doc, "原始数据保持不变", "初始化历史 sales.csv 只读。网页新增的真实销售记录单独写入 runtime_actual_sales.csv；推荐时再合并使用。")
    page_break(doc)

    add_heading(doc, "九、当前评估结果", 1)
    add_body(doc, "以下结果来自历史回放：系统在不知道当天真实销量的前提下生成菜单，再与历史真实数据对比。它能说明预测和选菜能力，但不能直接证明上线后营业额增加。")
    add_heading(doc, "A湖畔冻结测试（2026-04-01 至 2026-06-10）", 2)
    add_kpi_strip(doc, [
        ("13.70%", "每餐总量平均误差"),
        ("50.65%", "命中菜逐菜平均误差"),
        ("69.11 万", "历史回放净收益代理"),
        ("0", "硬性搭配违规"),
    ])
    add_table(doc, ["补充指标", "结果", "含义"], [
        ["命中菜加权误差", "26.22%", "大销量菜权重更高，更接近整体备餐偏差"],
        ["营业额覆盖率", "60.46%", "推荐菜覆盖了历史真实营业额的比例"],
        ["菜品命中率", "60.43%", "推荐菜中，在历史真实菜单出现并售出的比例"],
        ["菜单填充率", "99.90%", "大多数餐次能够达到计划菜单规模"],
        ["节假日菜品数偏差", "+5.04 道/餐", "节假日仍有偏多推荐现象"],
    ], [2600, 1600, 5160], numeric_cols=(1,))
    add_heading(doc, "F历史到基础园跨食堂测试（2026-01-01 至 2026-06-10）", 2)
    add_kpi_strip(doc, [
        ("10.21%", "每餐总量平均误差"),
        ("56.88%", "命中菜逐菜平均误差"),
        ("152.11 万", "历史回放净收益代理"),
        ("0", "硬性搭配违规"),
    ])
    add_table(doc, ["补充指标", "结果", "含义"], [
        ["命中菜加权误差", "21.58%", "整体备餐误差低于逐菜平均误差"],
        ["营业额覆盖率", "66.62%", "推荐菜覆盖的历史真实营业额比例"],
        ["菜品命中率", "62.42%", "推荐菜在历史真实菜单中的命中比例"],
        ["冷启动前7天总量误差", "20.15%", "刚切换到新食堂时误差较高，30天后降到8.93%"],
        ["相邻菜单重复率", "91.70%", "高频策略稳定但重复较高，潜力菜模块用于缓解"],
    ], [2600, 1600, 5160], numeric_cols=(1,))
    page_break(doc)

    add_heading(doc, "十、当前优势", 1)
    add_table(doc, ["优势", "说明"], [
        ["可解释", "每道菜可说明近期表现、长期表现、出菜频率、稳定性和搭配原因"],
        ["可控", "工作人员决定最终菜单；规则、参数、历史和输出都可追溯"],
        ["不依赖外网", "业务数据在本机保存和处理，降低数据外发风险"],
        ["防止未来数据混入", "预测日只使用此前销售历史；当天实际销量不会提前进入"],
        ["支持特殊日期", "周末、调休、寒暑假、法定节假日和天气可作为补充信息"],
        ["兼顾稳定与探索", "常用菜保证稳定，潜力菜提供新选择且不破坏原策略"],
        ["数据更新较安全", "预览、确认、去重、备份、失败恢复和文件锁均已具备"],
        ["跨平台", "Windows、macOS、Linux 均可在本地运行"],
    ], [2100, 7260])
    add_callout(doc, "适合的使用方式", "把系统作为菜单制定前的参考工具：先看推荐和数量，再结合库存、人员、窗口和现场经验确认。", fill=PALE_GREEN, accent=GREEN)
    page_break(doc)

    add_heading(doc, "十一、当前缺点与风险", 1)
    add_table(doc, ["问题", "当前表现", "业务影响"], [
        ["极端假期识别不足", "寒假与春节重叠、校园客流接近极低时仍可能明显高估", "菜品数和备餐量可能过多"],
        ["逐菜误差仍高", "A湖畔命中菜平均误差50.65%，基础园56.88%", "具体菜品的建议量仍需人工修正"],
        ["高频策略重复率高", "基础园相邻同餐次菜单重复率91.70%", "菜单稳定，但多样性不足"],
        ["节假日菜单仍偏多", "A湖畔平均多5.04道/餐，基础园假期平均多10.97道/餐", "特殊日期仍可能增加生产负担"],
        ["冷启动有磨合期", "基础园前7天总量误差20.15%，30天后8.93%", "新食堂或新窗口刚接入时需更谨慎"],
        ["低频/新菜证据少", "低频菜样本少，新菜没有真实销售历史", "推荐理由可以成立，但销量量化不稳定"],
        ["依赖菜品标签质量", "类别、食材、做法或相似菜标签错误会影响搭配判断", "需要持续维护菜品库"],
        ["经营指标仍是代理", "损耗按过度预测营业额的50%估算，缺少真实成本与剩餐数据", "不能把净收益直接当成实际利润"],
        ["历史回放不是因果证明", "没有真实上线A/B对照", "不能声称系统带来了营业额增量"],
    ], [2050, 3650, 3660])
    add_heading(doc, "需要特别向甲方说明", 2)
    add_callout(doc, "不夸大效果", "当前结果说明系统能够较稳定地预测当餐总量并提供有依据的菜单建议；它还不能证明上线后一定增加营业额或减少多少浪费。真正效果需要通过一段时间的试运行记录来验证。", fill="FCEFEF", accent=RED)
    page_break(doc)

    add_heading(doc, "十二、建议的下一阶段", 1)
    add_table(doc, ["阶段", "建议动作", "期望结果"], [
        ["1. 小范围试运行", "选择1个食堂、1至2个餐次，运行4至8周；保留人工确认", "验证系统能否融入真实排菜流程"],
        ["2. 补充关键数据", "记录最终采用菜单、实际备餐、售出、剩餐、成本、人工调整原因", "把代理净收益替换为更真实的经营指标"],
        ["3. 完善特殊日期", "增加低客流/特殊营业状态或管理人员手动客流等级", "解决寒暑假和春节重叠时的极端高估"],
        ["4. 降低重复", "为高频菜单设置按周类别目标、轮换提醒和重复率看板", "在不损害畅销菜覆盖的前提下提高多样性"],
        ["5. 提升逐菜预测", "针对低频菜、晚餐和特殊日期分别校准；积累潜力菜试餐反馈", "降低具体菜品备餐偏差"],
        ["6. 建立运营看板", "按周查看总量误差、逐菜误差、采用率、剩餐率和人工调整", "让项目从一次性模型变成持续改进工具"],
    ], [1700, 4660, 3000])
    add_heading(doc, "建议验收口径", 2)
    for t in [
        "功能层面：推荐、筛选、缓冲区、搭配检查、保存和数据更新稳定可用。",
        "数据层面：目标日不使用未来销售，新增数据可追溯、可去重、可恢复。",
        "业务层面：连续试运行期间统计人工采用率、总量误差、逐菜加权误差和剩餐变化。",
        "管理层面：明确系统是辅助建议，最终菜单由食堂管理人员确认。",
    ]:
        add_bullet(doc, t, num_id)
    add_callout(doc, "建议决策", "批准进入“受控试运行 + 数据补充”阶段。先验证真实运营价值，再决定是否扩大到更多食堂和窗口。", fill=PALE_GREEN, accent=GREEN)
    page_break(doc)

    add_heading(doc, "附录：主要指标的通俗解释", 1)
    add_table(doc, ["指标", "怎么理解", "注意事项"], [
        ["每餐总量平均误差", "每餐预测总销量与真实总销量相差多少，百分比越低越好", "真实销量特别小时会被放大"],
        ["命中菜逐菜平均误差", "对推荐且真实售出的每道菜，分别计算误差后求平均", "低销量菜会让该指标波动较大"],
        ["命中菜加权误差", "把所有命中菜的绝对误差加总，再除以真实销量总和", "更能反映整体备餐偏差"],
        ["菜品命中率", "推荐菜中实际在历史菜单出现并售出的比例", "历史没出现不一定代表真实上线不能卖"],
        ["营业额覆盖率", "推荐菜覆盖了历史真实营业额的比例", "不是上线后的新增营业额"],
        ["净收益代理", "推荐菜历史实际营业额 - 0.5 × 过度预测营业额", "0.5是假设，不能等同实际利润"],
        ["菜单数量误差", "推荐菜品数量与历史真实菜品数量的差距", "需要结合现场产能和窗口安排"],
        ["硬性违规", "重复菜或明确冲突菜等必须避免的问题", "当前测试中为0"],
    ], [2100, 4500, 2760])
    add_heading(doc, "资料口径", 2)
    add_body(doc, "本报告基于当前冻结策略 factor_capacity_v1_trial18、项目根目录 README、smart_canteen_web/README、trial18_train_validation_test_summary.csv 以及 F→基础园跨食堂测试结果整理。数据截至2026年6月10日；报告日期为2026年8月24日。")
    add_body(doc, "报告中的“净收益”均为历史回放代理值，不代表相对原菜单增加的营业额，也不代表财务口径利润。")

    core = doc.core_properties
    core.title = "智慧食堂菜单推荐系统项目阶段报告"
    core.subject = "甲方项目汇报材料"
    core.author = "智慧食堂项目组"
    core.keywords = "智慧食堂, 菜单推荐, 项目总结, 试运行"
    core.comments = "根据当前冻结模型、网页系统和已完成测试整理"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
